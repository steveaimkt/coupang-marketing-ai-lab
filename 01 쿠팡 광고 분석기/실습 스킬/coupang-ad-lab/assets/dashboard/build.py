"""설정 파일을 읽어 원본 XLSX/CSV로 로컬 분석기를 만든다.

만들 때마다 결과물/{자료 날짜}_{만든 시각 HHMM}/ 새 폴더에 넣는다. 지난 결과물은 덮어쓰지 않는다.
자료 날짜는 읽은 데이터 날짜 폴더 이름(예 2026-10-20)이고, 시각은 이 컴퓨터의 현지 시각이다.
같은 이름이 있으면 _2, _3 을 붙인다."""
from pathlib import Path
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
import argparse
import shutil
import csv
import hashlib
import html
import json
import re
import yaml
import openpyxl



def apply_claude_md(root, config):
    """작업 폴더 CLAUDE.md 「판단 기준」 줄이 있으면 config 보다 앞선다. 기준은 CLAUDE.md 한 곳에 둔다."""
    import re as _re
    from pathlib import Path as _P
    p = _P(root) / 'CLAUDE.md'
    if not p.exists():
        return config
    keys = ('attribution_days', 'waste_cost_krw', 'star_roas_pct', 'potential_roas_pct',
            'minimum_orders_for_star', 'target_roas_pct', 'margin_pct')
    out = dict(config)
    for m in _re.finditer(r'^\s*[-*]?\s*`?(' + '|'.join(keys) + r')`?\s*:\s*([^\s#|]+)', p.read_text(encoding='utf-8'), _re.M):
        k, v = m.group(1), m.group(2).strip('`')
        out[k] = None if v.lower() in ('null', 'none', '없음') else (float(v) if '.' in v else int(v))
    return out


def resolve_input(root, spec):
    """input_file 이 *·? 를 품으면 데이터/ 안의 가장 최근 날짜 폴더에서 하나를 고른다.
    폴더 이름은 리포트 마지막 날짜(2026-10-20 꼴)라 이름순이 곧 날짜순이다."""
    from pathlib import Path as _P
    if not any(ch in str(spec) for ch in '*?['):
        p = _P(spec)
        return p if p.is_absolute() else root / p
    found = [m for m in sorted(root.glob(str(spec))) if not m.name.startswith(('~$', '.'))]
    if not found:
        raise FileNotFoundError(f'매출 성장 광고 보고서를 찾지 못했습니다: {spec} · 데이터/ 안에 날짜 폴더(예 2026-10-25)를 만들고 받은 파일을 이름 그대로 넣으세요.')
    latest = max(m.parent for m in found)
    pick = [m for m in found if m.parent == latest]
    if len(pick) > 1:
        raise ValueError(f'{latest.name} 폴더에 매출 성장 광고 보고서가 {len(pick)}개 있습니다. 하나만 남기세요: ' + ', '.join(m.name for m in pick))
    return pick[0]


def new_result_folder(base, source, last_day):
    """결과물/{자료 날짜}_{HHMM}/ 를 새로 만든다. 같은 이름이 있으면 _2, _3 을 붙인다. 있는 폴더는 건드리지 않는다."""
    day = source.parent.name if re.fullmatch(r'\d{4}-\d{2}-\d{2}', source.parent.name) else last_day
    stem = f"{day}_{datetime.now().strftime('%H%M')}"
    base.mkdir(parents=True, exist_ok=True)
    n = 1
    while True:
        folder = base / (stem if n == 1 else f'{stem}_{n}')
        try:
            folder.mkdir()
            return folder
        except FileExistsError:
            n += 1

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--config', type=Path, help='생성된 분석 스킬의 config.md')
args = parser.parse_args()
if args.config:
    SKILL = args.config.resolve().parent
    if SKILL.parent.name != 'skills' or SKILL.parent.parent.name != '.claude':
        parser.error('config.md는 작업폴더/.claude/skills/스킬이름 안에 있어야 합니다.')
    ROOT = SKILL.parent.parent.parent
else:
    ROOT = next((p for p in Path(__file__).resolve().parents
                 if (p / '.claude/skills/coupang-ad-dashboard/config.md').exists()), None)
    if ROOT is None:
        parser.error('--config로 실제 분석 스킬 설정을 지정하세요.')
    SKILL = ROOT / '.claude/skills/coupang-ad-dashboard'
CONFIG = yaml.safe_load(re.search(r'```yaml\n(.*?)```', (SKILL / 'config.md').read_text(), re.S)[1])
CONFIG = apply_claude_md(ROOT, CONFIG)
source = resolve_input(ROOT, CONFIG['input_file'])
results = Path(CONFIG['output_directory'])
if not results.is_absolute():
    results = ROOT / results
days = CONFIG['attribution_days']
assert days in (1, 14)
assert CONFIG['star_roas_pct'] >= CONFIG['potential_roas_pct'] >= 0
assert CONFIG['waste_cost_krw'] > 0 and CONFIG['minimum_orders_for_star'] >= 1
assert CONFIG['target_roas_pct'] is None or CONFIG['target_roas_pct'] >= 0
CONFIG.setdefault('margin_pct', None)
assert CONFIG['margin_pct'] is None or 0 < CONFIG['margin_pct'] <= 100, '마진율(margin_pct)은 0보다 크고 100 이하이거나 null 이어야 합니다'
before = hashlib.sha256(source.read_bytes()).hexdigest()
if source.suffix.lower() == '.xlsx':
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet_name = CONFIG.get('sheet_name')
    if not sheet_name and len(wb.sheetnames) == 1:
        sheet_name = wb.sheetnames[0]
    if sheet_name not in wb.sheetnames:
        raise ValueError(f'시트를 선택해 sheet_name에 저장하세요: {wb.sheetnames}')
    sheet = wb[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(x).strip() if x is not None else '' for x in next(iterator)]
    original = list(iterator)
    wb.close()
    origin_detail = f'시트 {sheet_name}'
elif source.suffix.lower() == '.csv':
    with source.open(encoding=CONFIG.get('csv_encoding', 'utf-8-sig'), newline='') as f:
        iterator = csv.reader(f, delimiter=CONFIG.get('csv_delimiter', ','))
        headers = [x.strip() for x in next(iterator)]
        original = list(iterator)
    origin_detail = f"CSV / {CONFIG.get('csv_encoding', 'utf-8-sig')}"
else:
    raise ValueError('지원 입력: XLSX 또는 CSV')
required = ['날짜', '캠페인 ID', '캠페인명', '광고 노출 지면', '키워드', '노출수', '클릭수', '광고비', f'총 주문수({days}일)', f'총 전환매출액({days}일)']
missing = [k for k in required if k not in headers]
assert not missing, f'필수 열 누락: {missing}'
group_header = '광고그룹' if '광고그룹' in headers else '광고그룹명'
assert group_header in headers, '광고그룹 열 누락'

def number(v, count=False):
    if v is None or str(v).strip() == '':
        raise ValueError('비어 있는 성과 값')
    n = Decimal(str(v).replace(',', '').strip())
    if not n.is_finite() or n < 0 or (count and n != n.to_integral()):
        raise ValueError(f'잘못된 성과 값: {v}')
    return float(n) if not count else int(n)

def get_date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    if re.fullmatch(r'\d{8}(\.0)?', s):
        return datetime.strptime(s[:8], '%Y%m%d').strftime('%Y-%m-%d')
    return date.fromisoformat(s[:10]).isoformat()

def label(v):
    return '' if v is None else str(v).strip()

rows, excluded = [], []
outside = 0
for line, raw in enumerate(original, 2):
    r = dict(zip(headers, raw))
    try:
        d = get_date(r['날짜'])
        if (CONFIG['date_start'] and d < str(CONFIG['date_start'])) or (CONFIG['date_end'] and d > str(CONFIG['date_end'])):
            outside += 1
            continue
        row = dict(d=d, cid=label(r['캠페인 ID']), c=label(r['캠페인명']), g=label(r[group_header]), p=label(r['광고 노출 지면']) or '지면 미지정', k=label(r['키워드']) or '-', i=number(r['노출수'], True), clicks=number(r['클릭수'], True), cost=number(r['광고비']), orders=number(r[f'총 주문수({days}일)'], True), revenue=number(r[f'총 전환매출액({days}일)']))
        if not row['cid']:
            raise ValueError('캠페인 ID 없음')
        rows.append(row)
    except (ValueError, InvalidOperation, TypeError) as e:
        excluded.append({'row': line, 'reason': str(e)})
assert rows, '분석할 유효 행이 없습니다.'
totals = {k: float(sum(Decimal(str(r[k])) for r in rows)) for k in ('i', 'clicks', 'cost', 'orders', 'revenue')}
meta = dict(filename=source.name, detail=origin_detail, start=min(r['d'] for r in rows), end=max(r['d'] for r in rows), rows=len(rows), original_rows=len(original), columns=len(headers), excluded=excluded, outside=outside, sha256=before)
safe_config = {k: v for k, v in CONFIG.items() if k not in ('input_file', 'output_directory')}
payload = json.dumps(dict(meta=meta, config=safe_config, rows=rows), ensure_ascii=False, separators=(',', ':'), default=str).replace('<', '\\u003c').replace('>', '\\u003e').replace('&', '\\u0026')
title = html.escape(CONFIG['dashboard_title'])
document = f'''<!doctype html>
<html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{title}</title><link rel="stylesheet" href="styles.css"><script id="report-data" type="application/json">{payload}</script><script src="script.js" defer></script></head>
<body><a class="skip" href="#main">분석 내용으로 이동</a>
<header class="topbar"><div class="brand"><span class="brandmark" aria-hidden="true">C</span><div><strong>{title}</strong><span>쿠팡 광고 · 성과 리포트</span></div></div><div class="topmeta"><span class="dot"></span> 광고 성과 리포트 <span id="period" class="period"></span></div></header>
<div class="app"><nav class="navigation" aria-label="광고 분석 메뉴"><div role="tablist" id="tabs" aria-label="분석 메뉴"></div></nav>
<main id="main" tabindex="-1"><div class="pageheading"><div><div class="eyebrow">캠페인 분석</div><h1 id="page-title">광고 성과를 한눈에</h1><p id="page-description"></p></div><div class="heading-controls"><label for="campaign">캠페인 범위</label><select id="campaign"><option value="">전체 캠페인</option></select></div></div>
<div id="source-context" class="source-context"></div><div id="panel" role="tabpanel"></div>
<footer><span id="footer-source"></span><span>광고 전환매출 기준 · 비율은 합계로 재계산</span></footer></main></div>
<noscript>이 분석기는 JavaScript가 켜진 브라우저에서 열어 주세요.</noscript></body></html>'''
# 화면 자산은 분석 스킬 assets/dashboard 에서 가져온다. 화면을 고치려면 그쪽을 고친다.
asset_dir = SKILL / 'assets/dashboard'
if not asset_dir.is_dir():
    asset_dir = Path(__file__).resolve().parent
for name in ('styles.css', 'script.js'):
    if not (asset_dir / name).is_file():
        raise FileNotFoundError(f'화면 자산 없음: {asset_dir / name}')
# 만들 때마다 새 폴더. 지난 결과물은 그대로 남는다.
out = new_result_folder(results, source, meta['end'])
for name in ('styles.css', 'script.js'):
    shutil.copy2(asset_dir / name, out / name)
(out / 'index.html').write_text(document, encoding='utf-8')
assert before == hashlib.sha256(source.read_bytes()).hexdigest(), '원본 파일이 변경됨'
report = f'''# 생성 확인 · {out.name}

- 결과물 폴더: {out.relative_to(ROOT) if out.is_relative_to(ROOT) else out}/
- 원본: {source.parent.name}/{source.name} / {origin_detail}
- SHA-256: {before}
- 원본 {len(original):,}행, {len(headers)}개 열 / 분석 {len(rows):,}행
- 기간: {meta['start']}~{meta['end']} / 기간 밖 {outside}행 / 오류로 보류 {len(excluded)}행
- 전환 기준: {days}일 / 목표 ROAS: {str(CONFIG['target_roas_pct'])+'%' if CONFIG['target_roas_pct'] is not None else '미설정'}
- 마진율: {str(CONFIG['margin_pct'])+'% / 손익분기 ROAS '+format(10000/CONFIG['margin_pct'], '.1f')+'%' if CONFIG['margin_pct'] is not None else '미설정 (손익분기 판정 숨김)'}
- 효자: ROAS {CONFIG['star_roas_pct']}% 이상 · 주문 {CONFIG['minimum_orders_for_star']}건 이상 / 낭비 검토: 광고비 {CONFIG['waste_cost_krw']:,}원 이상 · 주문 0건
- 집계: 광고비 {totals['cost']:,.0f}원, 광고 전환매출 {totals['revenue']:,.0f}원, 주문 {totals['orders']:,.0f}건
- 원본 수정 없음: 생성 전후 SHA-256 일치
- 날짜·필수 열·유효 성과 값 확인 완료. 잘못된 행을 0으로 대체하지 않음.
- 이번 실행의 원본 독립 대조·브라우저·반복 실행: 미확인. 실행한 검사만 이 문서에 추가한다.
- 지난 결과물은 결과물/ 안 다른 폴더에 그대로 있다. 지난 통과를 이번 실행의 통과로 보지 않는다.

## 보류 행

{json.dumps(excluded, ensure_ascii=False) if excluded else '없음'}
'''
(out / '검증결과.md').write_text(report, encoding='utf-8')
shown = out.relative_to(ROOT) if out.is_relative_to(ROOT) else out
print(f'{shown}/ 에 만들었습니다')
print(json.dumps(dict(output=str(out), folder=out.name, totals=totals, meta=meta), ensure_ascii=False))
