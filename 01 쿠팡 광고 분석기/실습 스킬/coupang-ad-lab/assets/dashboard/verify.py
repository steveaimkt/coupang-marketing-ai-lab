"""원본을 별도로 읽어 HTML 내장 행과 Decimal 합계를 검산한다. 화면 검사는 별도이다.

따로 지정이 없으면 가장 최근 결과물(결과물/ 안 폴더 이름순 마지막)을 맞춰 본다.
지난 결과물은 --dir 2026-10-20_1405 처럼 폴더를 지정한다."""
import argparse
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
import yaml




def apply_claude_md(root, config):
    """작업 폴더 CLAUDE.md 「판단 기준」 줄이 있으면 config 보다 앞선다. 기준은 CLAUDE.md 한 곳에 둔다."""
    import re as _re
    from pathlib import Path as _P
    p = _P(root) / 'CLAUDE.md'
    if not p.exists():
        return config
    keys = ('attribution_days', 'waste_cost_krw', 'star_roas_pct', 'potential_roas_pct',
            'minimum_orders_for_star', 'target_roas_pct', 'margin_pct')
    out = dict(config)
    for m in _re.finditer(r'^\s*[-*]?\s*`?(' + '|'.join(keys) + r')`?\s*:\s*([^\s#|]+)', p.read_text(encoding='utf-8'), _re.M):
        k, v = m.group(1), m.group(2).strip('`')
        out[k] = None if v.lower() in ('null', 'none', '없음') else (float(v) if '.' in v else int(v))
    return out


def resolve_input(root, spec):
    """input_file 이 *·? 를 품으면 데이터/ 안의 가장 최근 날짜 폴더에서 하나를 고른다.
    폴더 이름은 리포트 마지막 날짜(2026-10-20 꼴)라 이름순이 곧 날짜순이다."""
    from pathlib import Path as _P
    if not any(ch in str(spec) for ch in '*?['):
        p = _P(spec)
        return p if p.is_absolute() else root / p
    found = [m for m in sorted(root.glob(str(spec))) if not m.name.startswith(('~$', '.'))]
    if not found:
        raise FileNotFoundError(f'매출 성장 광고 보고서를 찾지 못했습니다: {spec} · 데이터/ 안에 날짜 폴더(예 2026-10-25)를 만들고 받은 파일을 이름 그대로 넣으세요.')
    latest = max(m.parent for m in found)
    pick = [m for m in found if m.parent == latest]
    if len(pick) > 1:
        raise ValueError(f'{latest.name} 폴더에 매출 성장 광고 보고서가 {len(pick)}개 있습니다. 하나만 남기세요: ' + ', '.join(m.name for m in pick))
    return pick[0]


RESULT_NAME = re.compile(r'^(\d{4}-\d{2}-\d{2}_\d{4})(?:_(\d+))?$')


def result_folders(base):
    """결과물/ 안의 {자료 날짜}_{HHMM}[_N] 폴더를 이름순으로 돌려준다. 배포용 같은 다른 폴더는 뺀다."""
    if not Path(base).is_dir():
        return []
    found = [d for d in Path(base).iterdir() if d.is_dir() and RESULT_NAME.match(d.name)]
    return sorted(found, key=lambda d: (RESULT_NAME.match(d.name)[1], int(RESULT_NAME.match(d.name)[2] or 1)))


def latest_result(base):
    """가장 최근 결과물 = 결과물/ 안 폴더 이름순 마지막."""
    folders = result_folders(base)
    if not folders:
        raise FileNotFoundError(f'결과물이 아직 없습니다: {base} · 먼저 분석기를 만드세요.')
    return folders[-1]


def candidates(root, spec):
    """input_file 규칙에 맞는 원본 후보 전부. 지난 결과물을 맞춰 볼 때 쓴다."""
    if not any(ch in str(spec) for ch in '*?['):
        p = Path(spec)
        return [p if p.is_absolute() else root / p]
    return [m for m in sorted(root.glob(str(spec))) if not m.name.startswith(('~$', '.'))]


def classify(r, config):
    """script.js classify 와 같은 순서를 원본 합계(Decimal)로 다시 쓴다. 화면 코드를 가져오지 않는다."""
    roas = r['revenue'] / r['cost'] * 100 if r['cost'] else None
    if r['k'] in ('-', ''):
        return '키워드 미지정'
    if roas is not None and roas >= config['star_roas_pct'] and r['orders'] >= config['minimum_orders_for_star']:
        return '효자'
    if r['cost'] >= config['waste_cost_krw'] and r['orders'] == 0:
        return '낭비 검토'
    if r['orders'] > 0 and roas is not None and roas < config['potential_roas_pct']:
        return '가능성 검토'
    if r['orders'] > 0:
        return '주문 발생'
    if r['cost'] > 0 or r['clicks'] > 0:
        return '관찰'
    return '활동 없음'


def judge(total, groups, config):
    """손익분기 · 키워드 분류 · 지면 점검을 원본에서 다시 센다. 화면(verify-browser full)이 이 값과 맞아야 한다."""
    margin = config.get('margin_pct')
    breakeven = Decimal(10000) / Decimal(str(margin)) if margin else None
    counts = {}
    for r in groups['keyword']:
        c = classify(r, config)
        counts[c] = counts.get(c, 0) + 1
    kw = [r for r in groups['keyword'] if r['k'] not in ('-', '')]
    cost_kw = sum((r['cost'] for r in kw), Decimal(0))
    orders_kw = sum((r['orders'] for r in kw), Decimal(0))
    out = dict(counts=counts, star=counts.get('효자', 0), waste=counts.get('낭비 검토', 0),
               keyword_cpa=cost_kw / orders_kw if orders_kw else None,
               unassigned_cost=sum((r['cost'] for r in groups['keyword'] if r['k'] in ('-', '')), Decimal(0)),
               margin_pct=margin, breakeven_roas_pct=breakeven)
    if breakeven is None:
        out.update(breakeven_status=None, profit=None, loss=None, placement_check=[], placement_check_cost=None)
        return out
    out['breakeven_status'] = None if total['roas'] is None else ('위' if total['roas'] >= breakeven else '아래')
    roas_of = lambda r: r['revenue'] / r['cost'] * 100 if r['cost'] else None
    out['profit'] = sum(1 for r in kw if roas_of(r) is not None and roas_of(r) >= breakeven)
    out['loss'] = sum(1 for r in kw if roas_of(r) is not None and roas_of(r) < breakeven)
    checks = sorted((r for r in groups['placement'] if roas_of(r) is not None and roas_of(r) < breakeven), key=lambda r: -r['cost'])
    out['placement_check'] = [r['p'] for r in checks]
    out['placement_check_cost'] = sum((r['cost'] for r in checks), Decimal(0))
    return out


def judgement_line(j):
    parts = [f"효자 {j['star']}개", f"낭비 검토 {j['waste']}개"]
    if j['keyword_cpa'] is not None:
        parts.append(f"키워드 평균 CPA {j['keyword_cpa']:,.0f}원")
    if j['breakeven_roas_pct'] is None:
        parts.append('마진율 미설정 · 손익분기 판정 없음')
    else:
        parts.append(f"손익분기 ROAS {j['breakeven_roas_pct']:.1f}% · 전체 ROAS 손익분기 {j['breakeven_status']}")
        parts.append(f"이익 키워드 {j['profit']}개 · 손해 키워드 {j['loss']}개")
        parts.append(f"지면 점검 {len(j['placement_check'])}곳" + (f" ({', '.join(j['placement_check'])} · 광고비 {j['placement_check_cost']:,.0f}원)" if j['placement_check'] else ''))
    return ' / '.join(parts)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--config', type=Path, required=True)
    p.add_argument('--dir', help='맞춰 볼 결과물 폴더. 비우면 가장 최근 결과물')
    args = p.parse_args()
    config_path = args.config.resolve()
    config = yaml.safe_load(re.search(r'```yaml\n(.*?)```', config_path.read_text(), re.S)[1])
    root = config_path.parents[3]
    config = apply_claude_md(root, config)
    config.setdefault('margin_pct', None)
    results = Path(config['output_directory'])
    results = results if results.is_absolute() else root / results
    if args.dir:
        given = Path(args.dir)
        out = next((d for d in (given, root / given, results / given) if (d / 'index.html').is_file()), None)
        assert out is not None, f'결과물 폴더를 찾지 못했습니다: {args.dir}'
    else:
        out = latest_result(results)
    out = out.resolve()
    html = (out / 'index.html').read_bytes()
    saved = json.loads(re.search(r'<script id="report-data" type="application/json">(.*?)</script>', html.decode(), re.S)[1])
    expected_config = {k: v for k, v in config.items() if k not in ('input_file', 'output_directory')}
    if args.dir:
        # 지난 결과물은 만들 때 읽은 원본과 그때의 기준으로 맞춰 본다.
        found = [m for m in candidates(root, config['input_file'])
                 if m.is_file() and hashlib.sha256(m.read_bytes()).hexdigest() == saved['meta']['sha256']]
        assert found, f"이 결과물을 만들 때 읽은 원본({saved['meta']['filename']})을 데이터/ 에서 찾지 못했습니다"
        source = found[0]
        config = {**config, **saved['config']}
    else:
        source = resolve_input(root, config['input_file'])
        assert hashlib.sha256(source.read_bytes()).hexdigest() == saved['meta']['sha256'], \
            f'가장 최근 결과물({out.name})이 지금 자료로 만든 것이 아닙니다. 다시 분석기를 만드세요'
        assert json.loads(json.dumps(expected_config, default=str)) == saved['config'], '설정과 화면이 일치하지 않습니다'
    sha = hashlib.sha256(source.read_bytes()).hexdigest()
    if source.suffix.lower() == '.xlsx':
        import openpyxl
        book = openpyxl.load_workbook(source, read_only=True, data_only=True)
        sheet = config.get('sheet_name') or (book.sheetnames[0] if len(book.sheetnames) == 1 else None)
        assert sheet in book.sheetnames, 'sheet_name을 지정하세요'
        records = list(book[sheet].values)
        book.close()
    else:
        with source.open(encoding=config.get('csv_encoding', 'utf-8-sig'), newline='') as f:
            records = list(csv.reader(f, delimiter=config.get('csv_delimiter', ',')))
    header = [str(c).strip() if c is not None else '' for c in records[0]]
    group_column = '광고그룹' if '광고그룹' in header else '광고그룹명'
    days = config['attribution_days']
    columns = dict(cost='광고비', revenue=f'총 전환매출액({days}일)', orders=f'총 주문수({days}일)', clicks='클릭수', i='노출수')
    text_columns = dict(cid='캠페인 ID', c='캠페인명', g=group_column, p='광고 노출 지면', k='키워드')
    assert set(['날짜', *columns.values(), *text_columns.values()]) <= set(header), '필수 열 누락'
    rows, excluded, outside = [], [], 0
    row_numbers = []
    for n, values in enumerate(records[1:], 2):
        raw = dict(zip(header, values))
        try:
            value = raw['날짜']
            if isinstance(value, (datetime, date)):
                d = value.strftime('%Y-%m-%d')
            else:
                text = str(value).strip()
                d = (datetime.strptime(text[:8], '%Y%m%d').date().isoformat()
                     if re.fullmatch(r'\d{8}(\.0)?', text) else date.fromisoformat(text[:10]).isoformat())
            if config.get('date_start') and d < str(config['date_start']) or config.get('date_end') and d > str(config['date_end']):
                outside += 1
                continue
            row = {'d': d}
            for key, col in columns.items():
                val = Decimal(str(raw[col]).replace(',', '').strip())
                if not val.is_finite() or val < 0 or key in ('orders', 'clicks', 'i') and val != val.to_integral():
                    raise ValueError(col)
                row[key] = val
            for key, col in text_columns.items():
                row[key] = '' if raw.get(col) is None else str(raw[col]).strip()
            row['p'] = row['p'] or '지면 미지정'
            row['k'] = row['k'] or '-'
            if not row['cid']:
                raise ValueError('캠페인 ID 없음')
            rows.append(row)
            row_numbers.append(n)
        except (ValueError, TypeError, InvalidOperation):
            excluded.append(n)
    assert len(rows) == len(saved['rows'])
    assert excluded == [r['row'] for r in saved['meta']['excluded']]
    assert outside == saved['meta']['outside']
    for original, embedded in zip(rows, saved['rows']):
        for key in original:
            actual = Decimal(str(embedded[key])) if key in columns else embedded[key]
            assert actual == original[key], (key, actual, original[key])
    total = {key: sum((r[key] for r in rows), Decimal(0)) for key in columns}
    for key, num, den, mult in [('roas', 'revenue', 'cost', 100), ('ctr', 'clicks', 'i', 100),
                               ('cpc', 'cost', 'clicks', 1), ('cvr', 'orders', 'clicks', 100), ('cpa', 'cost', 'orders', 1)]:
        total[key] = total[num] / total[den] * mult if total[den] else None
    groups = {}
    for name, keys in [('daily', ['d']), ('campaign', ['cid']), ('placement', ['p']), ('keyword', ['cid', 'g', 'p', 'k'])]:
        buckets = {}
        for row, n in zip(rows, row_numbers):
            key = tuple(row[k] for k in keys)
            entry = buckets.setdefault(key, {**{k: row[k] for k in keys}, **{k: Decimal(0) for k in columns}, 'rownums': []})
            for col in columns:
                entry[col] += row[col]
            entry['rownums'].append(n)
        groups[name] = list(buckets.values())
        for col in columns:
            assert sum(r[col] for r in groups[name]) == total[col]
    assert sha == hashlib.sha256(source.read_bytes()).hexdigest()
    judgement = judge(total, groups, config)
    evidence = dict(total=total, groups=groups, sha256=sha, html_sha256=hashlib.sha256(html).hexdigest(),
                    config=saved['config'], excluded_rows=excluded, judgement=judgement)
    (out / '검산.json').write_text(json.dumps(evidence, ensure_ascii=False, default=lambda v: float(v) if isinstance(v, Decimal) else str(v)))
    with (out / '검증결과.md').open('a') as f:
        f.write('\n## 원본 독립 검산\n\n- 생성 코드의 파서·집계 함수를 가져오지 않고 원본을 다시 읽어 Decimal로 검산.\n'
                f'- 원본: {source.parent.name}/{source.name}\n'
                f'- 원본 SHA-256: {sha}\n- HTML SHA-256: {evidence["html_sha256"]}\n'
                f'- 유효 {len(rows)}행의 필드·합계 일치. 보류 {len(excluded)}행, 기간 밖 {outside}행 대조 일치.\n'
                '- 일별·캠페인·지면·키워드 합계와 원본 행 번호: 검산.json 참조.\n'
                f'- 판정 다시 세기 (전체 캠페인): {judgement_line(judgement)}\n'
                '- 브라우저 표시·상호작용: 이 검사만으로는 미확인.\n')
    shown = out.relative_to(root) if out.is_relative_to(root) else out
    print(json.dumps({'matched': True, 'dir': str(shown), 'rows': len(rows), 'total': total, 'judgement': judgement}, ensure_ascii=False, default=str))


if __name__ == '__main__':
    main()
